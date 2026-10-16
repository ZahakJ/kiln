#!/usr/bin/env python3
# SPDX-License-Identifier: AGPL-3.0-or-later
"""A second, independent reader for the feature fixtures exceljs cannot open.

    python3 -I crates/kiln-io-xlsx/tests/oracle/openpyxl-xlsx.py fixtures/xlsx/features/<name>.xlsx ...

writes fixtures/xlsx/read/<name>.openpyxl.json next to the exceljs files: per sheet, every cell's
constant value or formula, the merged ranges, the comments (author and text) and the tables
(name and range). Run with openpyxl 3.1 (`pip install openpyxl==3.1.5`); the JSON is committed,
so the test (tests/openpyxl_oracle.rs) needs no Python.
"""

import datetime
import json
import pathlib
import sys

import openpyxl


def serial(value, epoch):
    if isinstance(value, datetime.datetime):
        delta = value - datetime.datetime.combine(epoch, datetime.time())
    elif isinstance(value, datetime.date):
        delta = datetime.datetime.combine(value, datetime.time()) - datetime.datetime.combine(
            epoch, datetime.time()
        )
    elif isinstance(value, datetime.time):
        return (value.hour * 3600 + value.minute * 60 + value.second) / 86400
    else:
        return None
    days = delta.days + delta.seconds / 86400 + delta.microseconds / 86400e6
    # Excel's 1900 calendar counts the fictitious 1900-02-29.
    if epoch.year == 1899 and days < 60:
        days -= 1
    return days


def cell_json(cell, epoch):
    v = cell.value
    if v is None:
        return None
    if cell.data_type == "f":
        text = v if isinstance(v, str) else getattr(v, "text", str(v))
        return {"f": text.lstrip("=")}
    if cell.data_type == "e":
        return {"e": str(v)}
    if isinstance(v, bool):
        return {"b": v}
    if isinstance(v, (int, float)):
        return {"n": float(v)}
    s = serial(v, epoch)
    if s is not None:
        return {"n": s}
    return {"s": str(v)}


def main(paths):
    for p in map(pathlib.Path, paths):
        wb = openpyxl.load_workbook(p)
        epoch = (wb.epoch.date() if isinstance(wb.epoch, datetime.datetime) else wb.epoch)
        sheets = []
        for ws in wb.worksheets:
            cells = {}
            comments = []
            for row in ws.iter_rows():
                for c in row:
                    j = cell_json(c, epoch)
                    if j is not None:
                        cells[c.coordinate] = j
                    if c.comment is not None:
                        comments.append(
                            {"ref": c.coordinate, "author": c.comment.author or "", "text": c.comment.text}
                        )
            sheets.append(
                {
                    "name": ws.title,
                    "cells": cells,
                    "merges": sorted(str(m) for m in ws.merged_cells.ranges),
                    "comments": comments,
                    "tables": sorted(
                        ({"name": t.name, "ref": t.ref} for t in ws.tables.values()),
                        key=lambda t: t["name"],
                    ),
                }
            )
        out = p.parent.parent / "read" / (p.stem + ".openpyxl.json")
        out.write_text(
            json.dumps({"openpyxl": openpyxl.__version__, "sheets": sheets}, indent=1, ensure_ascii=False)
            + "\n"
        )
        print(out)


if __name__ == "__main__":
    main(sys.argv[1:])
