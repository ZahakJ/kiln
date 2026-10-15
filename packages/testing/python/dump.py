"""Dump a .docx/.xlsx/.pptx file as JSON using python-docx/openpyxl/python-pptx (test oracle).

    python -I dump.py <file>
"""

import datetime as dt
import json
import sys


def plain(v):
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, dt.timedelta):
        return v.total_seconds()
    if v is None or isinstance(v, (bool, int, float, str)):
        return v
    return str(v)


def dump_docx(path):
    import docx

    d = docx.Document(path)
    cp = d.core_properties
    return {
        "format": "docx",
        "paragraphs": [
            {
                "style": p.style.name if p.style is not None else None,
                "text": p.text,
                "runs": [{"text": r.text, "bold": r.bold, "italic": r.italic} for r in p.runs],
            }
            for p in d.paragraphs
        ],
        "tables": [[[c.text for c in row.cells] for row in t.rows] for t in d.tables],
        "sections": [
            {
                "orientation": str(s.orientation).split(".")[-1].split(" ")[0],
                "page_width": s.page_width,
                "page_height": s.page_height,
            }
            for s in d.sections
        ],
        "core": {k: plain(getattr(cp, k)) for k in ("title", "subject", "author", "keywords", "category", "revision", "created", "modified")},
    }


def dump_xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False, rich_text=False)
    sheets = []
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                v = c.value
                if hasattr(v, "text") and not isinstance(v, str):  # ArrayFormula and friends
                    v = getattr(v, "text", str(v))
                cells[c.coordinate] = {"value": plain(v), "number_format": c.number_format, "data_type": c.data_type}
        sheets.append({"name": ws.title, "state": ws.sheet_state, "dimensions": ws.dimensions, "cells": cells, "merged": sorted(str(r) for r in ws.merged_cells.ranges)})
    return {
        "format": "xlsx",
        "sheets": sheets,
        "defined_names": sorted(wb.defined_names.keys()),
        "epoch": wb.epoch.isoformat(),
    }


def dump_pptx(path):
    import pptx

    prs = pptx.Presentation(path)
    slides = []
    for s in prs.slides:
        shapes = []
        for sh in s.shapes:
            shapes.append(
                {
                    "name": sh.name,
                    "shape_type": str(sh.shape_type).split(" ")[0] if sh.shape_type is not None else None,
                    "text": sh.text_frame.text if sh.has_text_frame else None,
                    "left": sh.left,
                    "top": sh.top,
                    "width": sh.width,
                    "height": sh.height,
                }
            )
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide and s.notes_slide.notes_text_frame is not None else None
        slides.append({"layout": s.slide_layout.name, "hidden": s._element.get("show") == "0", "shapes": shapes, "notes": notes})
    return {"format": "pptx", "slide_width": prs.slide_width, "slide_height": prs.slide_height, "slides": slides}


def main():
    path = sys.argv[1]
    ext = path.rsplit(".", 1)[-1].lower()
    fn = {"docx": dump_docx, "xlsx": dump_xlsx, "pptx": dump_pptx}[ext]
    json.dump(fn(path), sys.stdout, ensure_ascii=False)


if __name__ == "__main__":
    main()
