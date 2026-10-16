#!/usr/bin/env python3
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Generates the chart fixtures of fixtures/xlsx/charts/ (WP-73): one workbook per chart type.

    python3 -m venv /tmp/venv && /tmp/venv/bin/pip install openpyxl==3.1.5
    /tmp/venv/bin/python fixtures/xlsx/charts/generate.py

openpyxl writes the classic DrawingML charts. It cannot write chartex parts, so treemap.xlsx and
sunburst.xlsx are written by openpyxl and then given a cx:chartSpace part, its drawing and the
relationships, exactly as Excel 2016 lays them out. Output is deterministic (fixed timestamps).
"""
import datetime
import io
import os
import zipfile

from openpyxl import Workbook
from openpyxl.chart import (AreaChart, BarChart, BubbleChart, DoughnutChart, LineChart, PieChart,
                            RadarChart, Reference, ScatterChart, Series)
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.trendline import Trendline

HERE = os.path.dirname(os.path.abspath(__file__))
FIXED = datetime.datetime(2026, 10, 16, 12, 0, 0)

QUARTERS = [("Quarter", "Materials", "Firing", "Events"),
            ("Q1", 730, 1180, 0), ("Q2", 505, 640, 0), ("Q3", 745, 1240, 120), ("Q4", 650, 1320, 380)]


def book(rows, title="Data"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    return wb, ws


def save(wb, name):
    buf = io.BytesIO()
    wb.save(buf)
    write_deterministic(buf.getvalue(), os.path.join(HERE, name))


def write_deterministic(data, path, extra=None):
    """Rewrites a zip with fixed timestamps and sorted-as-written entries, adding `extra` parts."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            body = src.read(info.filename)
            if extra and info.filename in extra:
                body = extra.pop(info.filename)
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, body)
        for name, body in (extra or {}).items():
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, body)
    with open(path, "wb") as f:
        f.write(out.getvalue())


def quarter_refs(ws):
    data = Reference(ws, min_col=2, min_row=1, max_col=4, max_row=5)
    cats = Reference(ws, min_col=1, min_row=2, max_row=5)
    return data, cats


def column_and_bar(direction, name):
    wb, ws = book(QUARTERS)
    for i, (grouping, overlap, title) in enumerate(
            [("clustered", None, "Clustered"), ("stacked", 100, "Stacked"),
             ("percentStacked", 100, "100% stacked")]):
        ch = BarChart()
        ch.type = direction
        ch.grouping = grouping
        if overlap is not None:
            ch.overlap = overlap
        ch.title = f"{title} {'column' if direction == 'col' else 'bar'}"
        ch.y_axis.title = "Spend (£)"
        ch.x_axis.title = "Quarter"
        ch.y_axis.majorGridlines = ch.y_axis.majorGridlines
        data, cats = quarter_refs(ws)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        if i == 0:
            ch.dataLabels = DataLabelList()
            ch.dataLabels.showVal = True
            ch.y_axis.number_format = '"£"#,##0'
            ch.legend.position = "b"
        ws.add_chart(ch, f"F{2 + i * 16}")
    save(wb, name)


def line():
    wb, ws = book(QUARTERS)
    for i, (title, markers, smooth) in enumerate(
            [("Line", False, False), ("Line with markers", True, False), ("Smooth line", True, True)]):
        ch = LineChart()
        ch.title = title
        data, cats = quarter_refs(ws)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        for s in ch.series:
            s.smooth = smooth
            if markers:
                s.marker.symbol = "circle"
                s.marker.size = 7
            else:
                s.marker.symbol = "none"
        if i == 2:
            ch.series[1].trendline = Trendline(trendlineType="linear", dispRSqr=True)
            ch.series[0].trendline = Trendline(trendlineType="movingAvg", period=2)
            ch.legend.position = "t"
        ws.add_chart(ch, f"F{2 + i * 16}")
    save(wb, "line.xlsx")


def area():
    wb, ws = book(QUARTERS)
    for i, grouping in enumerate(["standard", "stacked", "percentStacked"]):
        ch = AreaChart()
        ch.grouping = grouping
        ch.title = f"Area ({grouping})"
        data, cats = quarter_refs(ws)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        ws.add_chart(ch, f"F{2 + i * 16}")
    save(wb, "area.xlsx")


SHARES = [("Item", "Spend"), ("Gas and kiln time", 42), ("Clay and glaze", 24), ("Bursaries", 19),
          ("Shelves and repairs", 15)]


def pie():
    wb, ws = book(SHARES)
    ch = PieChart()
    ch.title = "Where the money went"
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showPercent = True
    ch.series[0].explosion = 0
    from openpyxl.chart.marker import DataPoint
    pt = DataPoint(idx=0, explosion=12)
    ch.series[0].dPt.append(pt)
    ws.add_chart(ch, "D2")
    save(wb, "pie.xlsx")


def doughnut():
    wb, ws = book(SHARES)
    ch = DoughnutChart()
    ch.title = "Ticket money"
    ch.holeSize = 55
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ws.add_chart(ch, "D2")
    save(wb, "doughnut.xlsx")


FIRINGS = [("Temperature (°C)", "Shrinkage (%)", "Absorption (%)", "Pieces")]
for i, t in enumerate(range(1000, 1310, 30)):
    FIRINGS.append((t, round(2 + (t - 1000) * 0.03 + (i % 3) * 0.4, 2), round(14 - (t - 1000) * 0.04 + (i % 2) * 0.5, 2), 8 + (i * 7) % 23))


def scatter():
    wb, ws = book(FIRINGS)
    for i, (title, style, lines) in enumerate([("Markers", "marker", False), ("Lines and markers", "lineMarker", True)]):
        ch = ScatterChart()
        ch.title = f"Firing curve: {title.lower()}"
        ch.style = 13
        ch.scatterStyle = style
        ch.x_axis.title = "Temperature (°C)"
        ch.y_axis.title = "Percent"
        ch.x_axis.scaling.min = 1000
        ch.x_axis.scaling.max = 1300
        ch.x_axis.majorUnit = 50
        xs = Reference(ws, min_col=1, min_row=2, max_row=len(FIRINGS))
        for col in (2, 3):
            s = Series(Reference(ws, min_col=col, min_row=1, max_row=len(FIRINGS)), xs, title_from_data=True)
            if not lines:
                s.graphicalProperties.line.noFill = True
                s.marker.symbol = "circle"
            ch.series.append(s)
        if i == 1:
            ch.series[0].trendline = Trendline(trendlineType="exp", dispEq=True)
        ws.add_chart(ch, f"F{2 + i * 16}")
    save(wb, "scatter.xlsx")


def bubble():
    wb, ws = book(FIRINGS)
    ch = BubbleChart()
    ch.title = "Pieces per firing"
    ch.style = 18
    xs = Reference(ws, min_col=1, min_row=2, max_row=len(FIRINGS))
    ys = Reference(ws, min_col=2, min_row=2, max_row=len(FIRINGS))
    sizes = Reference(ws, min_col=4, min_row=2, max_row=len(FIRINGS))
    s = Series(values=ys, xvalues=xs, zvalues=sizes, title="Firings")
    ch.series.append(s)
    ws.add_chart(ch, "F2")
    save(wb, "bubble.xlsx")


SKILLS = [("Skill", "Ana", "Ben"), ("Throwing", 9, 5), ("Trimming", 7, 6), ("Glazing", 6, 9),
          ("Loading", 8, 7), ("Firing", 5, 8), ("Teaching", 7, 4)]


def radar():
    wb, ws = book(SKILLS)
    for i, style in enumerate(["standard", "marker", "filled"]):
        ch = RadarChart()
        ch.type = style
        ch.title = f"Radar ({style})"
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_col=3, max_row=7), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=7))
        if style == "filled":
            ch.y_axis.delete = True
        ws.add_chart(ch, f"E{2 + i * 16}")
    save(wb, "radar.xlsx")


def combo():
    rows = [("Month", "Spend", "Members")] + [(m, s, n) for m, s, n in zip(
        ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"],
        [610, 640, 660, 580, 560, 565, 640, 650, 695, 720, 820, 810],
        [18, 18, 18, 19, 19, 19, 21, 21, 21, 23, 24, 24])]
    wb, ws = book(rows)
    bar = BarChart()
    bar.title = "Spend and members"
    bar.add_data(Reference(ws, min_col=2, min_row=1, max_row=13), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=13))
    bar.y_axis.title = "Spend (£)"
    bar.y_axis.majorGridlines = None
    ln = LineChart()
    ln.add_data(Reference(ws, min_col=3, min_row=1, max_row=13), titles_from_data=True)
    ln.y_axis.axId = 200
    ln.y_axis.title = "Members"
    ln.y_axis.crosses = "max"
    ln.series[0].marker.symbol = "circle"
    bar += ln
    bar.legend.position = "b"
    ws.add_chart(bar, "E2")
    save(wb, "combo.xlsx")


def axes():
    """Axis options: dates, log scale, reversed, fixed bounds, units, number formats."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Axes"
    ws.append(("Date", "Visitors", "Revenue"))
    start = datetime.date(2026, 1, 1)
    for i in range(12):
        ws.append((start + datetime.timedelta(days=30 * i), 10 ** (1 + i / 3), 120 + 35 * i))
    for r in range(2, 14):
        ws.cell(r, 1).number_format = "mmm yy"
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    ch = LineChart()
    ch.title = "Visitors (log scale)"
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=13), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=13))
    ch.y_axis.scaling.logBase = 10
    ch.y_axis.number_format = "#,##0"
    ch.x_axis.number_format = "mmm yy"
    ch.x_axis.majorTimeUnit = "months"
    from openpyxl.chart.axis import DateAxis
    ch.x_axis = DateAxis(crossAx=100)
    ch.x_axis.number_format = "mmm"
    ch.x_axis.majorTimeUnit = "months"
    ch.x_axis.title = "Month"
    ch.legend = None
    ws.add_chart(ch, "E2")
    ch2 = BarChart()
    ch2.type = "bar"
    ch2.title = "Revenue (reversed, fixed bounds)"
    ch2.add_data(Reference(ws, min_col=3, min_row=1, max_row=13), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=1, min_row=2, max_row=13))
    ch2.x_axis.scaling.orientation = "maxMin"
    ch2.y_axis.scaling.min = 0
    ch2.y_axis.scaling.max = 600
    ch2.y_axis.majorUnit = 100
    ch2.y_axis.number_format = '"£"0'
    ch2.legend.position = "l"
    ch2.layout = Layout(manualLayout=ManualLayout(x=0.25, y=0.2, h=0.6, w=0.7))
    ws.add_chart(ch2, "E18")
    save(wb, "axes.xlsx")


HIER = [("Region", "Studio", "Members"),
        ("North", "Leeds", 24), ("North", "York", 13), ("North", "Hull", 9),
        ("South", "Bristol", 31), ("South", "Exeter", 12),
        ("East", "Norwich", 17), ("East", "Ipswich", 8), ("East", "Cambridge", 21),
        ("West", "Hereford", 6)]

CHARTEX_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:cx="http://schemas.microsoft.com/office/drawing/2014/chartex"'


def chartex_xml(layout, title):
    n = len(HIER) - 1
    leaf = "".join(f'<cx:pt idx="{i}">{r[1]}</cx:pt>' for i, r in enumerate(HIER[1:]))
    parent = "".join(f'<cx:pt idx="{i}">{r[0]}</cx:pt>' for i, r in enumerate(HIER[1:]))
    vals = "".join(f'<cx:pt idx="{i}">{r[2]}</cx:pt>' for i, r in enumerate(HIER[1:]))
    layout_pr = '<cx:layoutPr><cx:parentLabelLayout val="overlapping"/></cx:layoutPr>' if layout == "treemap" else ""
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<cx:chartSpace {CHARTEX_NS}>'
            f'<cx:chartData><cx:data id="0"><cx:strDim type="cat"><cx:f>Data!$A$2:$B${n + 1}</cx:f>'
            f'<cx:lvl ptCount="{n}">{leaf}</cx:lvl><cx:lvl ptCount="{n}">{parent}</cx:lvl></cx:strDim>'
            f'<cx:numDim type="size"><cx:f>Data!$C$2:$C${n + 1}</cx:f><cx:lvl ptCount="{n}" formatCode="General">{vals}</cx:lvl></cx:numDim>'
            f'</cx:data></cx:chartData><cx:chart><cx:title pos="t" align="ctr" overlay="0"><cx:tx><cx:txData><cx:v>{title}</cx:v></cx:txData></cx:tx></cx:title>'
            f'<cx:plotArea><cx:plotAreaRegion><cx:series layoutId="{layout}" uniqueId="{{6A1E2B9C-0000-4000-8000-00000000000{1 if layout == "treemap" else 2}}}">'
            f'<cx:tx><cx:txData><cx:f>Data!$C$1</cx:f><cx:v>Members</cx:v></cx:txData></cx:tx>'
            f'<cx:dataLabels pos="{"inEnd" if layout == "treemap" else "ctr"}"><cx:visibility seriesName="0" categoryName="1" value="0"/></cx:dataLabels>'
            f'<cx:dataId val="0"/>{layout_pr}</cx:series></cx:plotAreaRegion></cx:plotArea>'
            f'<cx:legend pos="t" align="ctr" overlay="0"/></cx:chart></cx:chartSpace>')


DRAWING = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
           '<xdr:twoCellAnchor><xdr:from><xdr:col>4</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>1</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
           '<xdr:to><xdr:col>12</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>18</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
           '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><mc:Choice xmlns:cx1="http://schemas.microsoft.com/office/drawing/2015/9/8/chartex" Requires="cx1">'
           '<xdr:graphicFrame macro=""><xdr:nvGraphicFramePr><xdr:cNvPr id="2" name="Chart 1"/><xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr>'
           '<xdr:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></xdr:xfrm><a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/drawing/2014/chartex">'
           '<cx:chart xmlns:cx="http://schemas.microsoft.com/office/drawing/2014/chartex" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/>'
           '</a:graphicData></a:graphic></xdr:graphicFrame></mc:Choice><mc:Fallback><xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="2" name="Chart 1"/><xdr:cNvSpPr><a:spLocks noTextEdit="1"/></xdr:cNvSpPr></xdr:nvSpPr>'
           '<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="4876800" cy="3238500"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>'
           '<xdr:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US" sz="1100"/><a:t>This chart isn\'t available in your version of Excel.</a:t></a:r></a:p></xdr:txBody></xdr:sp></mc:Fallback></mc:AlternateContent>'
           '<xdr:clientData/></xdr:twoCellAnchor></xdr:wsDr>')


def hierarchy(layout, name, title):
    wb, ws = book(HIER)
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    ct = src.read("[Content_Types].xml").decode()
    ct = ct.replace("</Types>",
                    '<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
                    '<Override PartName="/xl/charts/chartEx1.xml" ContentType="application/vnd.ms-office.chartex+xml"/></Types>')
    sheet = src.read("xl/worksheets/sheet1.xml").decode()
    sheet = sheet.replace("</worksheet>", '<drawing r:id="rId1"/></worksheet>')
    if 'xmlns:r=' not in sheet.split(">", 2)[1]:
        sheet = sheet.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
    extra = {
        "[Content_Types].xml": ct.encode(),
        "xl/worksheets/sheet1.xml": sheet.encode(),
        "xl/worksheets/_rels/sheet1.xml.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/></Relationships>').encode(),
        "xl/drawings/drawing1.xml": DRAWING.encode(),
        "xl/drawings/_rels/drawing1.xml.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2014/relationships/chartEx" Target="../charts/chartEx1.xml"/></Relationships>').encode(),
        "xl/charts/chartEx1.xml": chartex_xml(layout, title).encode(),
    }
    write_deterministic(buf.getvalue(), os.path.join(HERE, name), extra)


if __name__ == "__main__":
    column_and_bar("col", "column.xlsx")
    column_and_bar("bar", "bar.xlsx")
    line()
    area()
    pie()
    doughnut()
    scatter()
    bubble()
    radar()
    combo()
    axes()
    hierarchy("treemap", "treemap.xlsx", "Members by studio")
    hierarchy("sunburst", "sunburst.xlsx", "Members by region")
